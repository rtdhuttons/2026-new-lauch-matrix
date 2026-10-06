"""TRM's chosen comparable projects, from the project sheet, for each project's
Investor tab.

Usage:
  python3 scripts/data/comparables-from-sheet.py exports/trm-map-projects-comparables.xlsx

Reads the "Projects" sheet (column O: comparable project(s), several split by
";" or " or "; column P: why it's comparable; column Q: distance limit), keeps
the projects still on the new launches map (matched by map project ID), and
for each comparable finds:
  - its address and position on OneMap (needs www.onemap.gov.sg): the result
    whose building name matches, nearest to the project
  - the straight-line distance to the project
  - tenure, units, developer and completion from the Huttons catalogue, when
    the comparable is a Huttons project
Writes src/features/selector/data/comparables/chosen.ts (don't edit by hand).
TRM's own words stay as written; the website's explanation of each choice is
in chosen-notes.ts.
"""

import json
import math
import re
import sys
import time
import urllib.parse
import urllib.request
from datetime import date
from pathlib import Path

import importlib.util

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "huttons"))
import client  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
CATALOGUE = ROOT / "src/features/catalogue/data/huttons-catalogue.ts"
OUT = ROOT / "src/features/selector/data/comparables/chosen.ts"
ONEMAP = "https://www.onemap.gov.sg/api/common/elastic/search?"
# Postal sector -> district, from scripts/data/sg-districts.py.
_spec = importlib.util.spec_from_file_location("sg_districts", Path(__file__).parent / "sg-districts.py")
_sg = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_sg)
DISTRICT_OF = _sg.DISTRICT_OF
CJK = re.compile(r"[⺀-鿿豈-﫿＀-￯]+")


def clean_name(name):
    return re.sub(r"\s+", " ", CJK.sub("", name or "")).strip()


UC = re.compile(r"\s*\(U/C\)", re.I)


def squash(s):
    """Names compared loosely: no "the", "(U/C)" or punctuation, "@" read as "at"."""
    s = UC.sub("", (s or "").lower()).replace("@", " at ")
    return re.sub(r"[^a-z0-9]", "", re.sub(r"^\s*the\s+", "", s))


def metres(a, b):
    p1, p2 = math.radians(a[0]), math.radians(b[0])
    h = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(b[1] - a[1]) / 2) ** 2
    return 2 * 6371000 * math.asin(math.sqrt(h))


def onemap(q):
    url = ONEMAP + urllib.parse.urlencode({"searchVal": q, "returnGeom": "Y", "getAddrDetails": "Y", "pageNum": 1})
    for _ in range(3):
        try:
            return json.load(urllib.request.urlopen(url, timeout=30)).get("results", [])
        except Exception:
            time.sleep(1)
    return []


ACRONYMS = {"MB", "BT", "GLS"}


def title(name):
    """"PASIR RIS 8" -> "Pasir Ris 8"; names already in mixed case are kept."""
    if name != name.upper():
        return name
    words = [w if w in ACRONYMS else w.capitalize() for w in name.split()]
    return " ".join(w.lower() if i and w in {"At", "Of", "Del"} else w for i, w in enumerate(words))


def locate(name, near):
    """OneMap's best match for a development's name: building name matches, nearest to `near`."""
    results = onemap(name)
    want = squash(name)
    exact = [r for r in results if squash(r.get("BUILDING")) == want]
    loose = [r for r in results if want in squash(r.get("BUILDING")) or squash(r.get("BUILDING")) in want and len(squash(r.get("BUILDING"))) > 4]
    pool = exact or loose
    if not pool:
        return None
    best = min(pool, key=lambda r: metres(near, (float(r["LATITUDE"]), float(r["LONGITUDE"]))))
    lat, lon = float(best["LATITUDE"]), float(best["LONGITUDE"])
    return {
        "address": best["ADDRESS"].replace(f" {best['BUILDING']}", "").title(),
        "underConstruction": bool(UC.search(best.get("BUILDING") or "")),
        "district": DISTRICT_OF.get(str(best.get("POSTAL") or "")[:2]),
        "lat": round(lat, 6),
        "lon": round(lon, 6),
        "km": round(metres(near, (lat, lon)) / 1000, 2),
        "exact": bool(exact),
    }


def main(path):
    t = CATALOGUE.read_text()
    fetched = re.search(r'fetched: "(\d{4}-\d{2}-\d{2})"', t).group(1)
    rows = [json.loads(line.strip().rstrip(",")) for line in t.splitlines() if line.strip().startswith('{"id"')]
    by_id = {r["id"]: r for r in rows}
    by_name = {squash(clean_name(r["name"])): r for r in rows}

    ws = openpyxl.load_workbook(path, data_only=True)["Projects"]
    head = [c.value for c in ws[5]]
    col = {name: head.index(name) for name in head if name}
    out, skipped = [], []
    for r in ws.iter_rows(min_row=7, values_only=True):
        pid, chosen = r[col["Map project ID"]], r[col["Comparable project(s) for resale and rental records"]]
        if not pid or not chosen or not str(chosen).strip():
            continue
        me = by_id.get(pid)
        if not me or not client.on_map(me, fetched):
            skipped.append(clean_name(r[col["Project"]]))
            continue
        near = (me["lat"], me["lon"])
        comps = []
        for name in re.split(r"\s*;\s*|\s+or\s+", str(chosen).strip(), flags=re.I):
            if not name:
                continue
            loc = locate(name, near)
            cat = by_name.get(squash(name))
            comps.append({
                "name": title(name.strip()),
                "sheetName": name.strip(),
                "address": loc["address"] if loc else None,
                "lat": loc["lat"] if loc else None,
                "lon": loc["lon"] if loc else None,
                "km": loc["km"] if loc else None,
                "located": ("onemap" if loc["exact"] else "onemap-partial") if loc else None,
                "underConstruction": loc["underConstruction"] if loc else None,
                "district": (cat["district"] if cat else None) or (loc["district"] if loc else None),
                "huttons": {
                    "tenure": cat["tenure"],
                    "totalUnits": cat["totalUnits"],
                    "developer": cat["developer"],
                    "completion": cat["completionDate"],
                    "district": cat["district"],
                } if cat else None,
            })
            time.sleep(0.2)
        limit = r[col["Distance limit (km, optional)"]]
        why = r[col["Why it's comparable (optional)"]]
        out.append({
            "slug": client.slugify(clean_name(me["name"])),
            "project": clean_name(me["name"]),
            "comparables": comps,
            "trmNote": str(why).strip() if why else None,
            "distanceLimitKm": float(limit) if limit not in (None, "") else None,
        })
        print(f"{out[-1]['project']}: " + "; ".join(f"{c['name']} ({c['km']} km{', partial match' if c['located'] == 'onemap-partial' else ''}{', not found' if c['located'] is None else ''}{', UNDER CONSTRUCTION' if c['underConstruction'] else ''})" for c in comps))
    body = json.dumps({"source": "TRM's project sheet", "checked": date.today().isoformat(), "catalogue": fetched, "projects": out}, indent=2, ensure_ascii=False)
    OUT.write_text(
        "// Generated by scripts/data/comparables-from-sheet.py from TRM's project sheet:\n"
        "// the comparable project(s) TRM chose for each project's resale and rental\n"
        "// records, located on OneMap. Don't edit by hand.\n\n"
        'import type { ChosenComparables } from "../../model/project";\n\n'
        f"export const chosenComparables: ChosenComparables = {body};\n"
    )
    print(f"{len(out)} projects written; not on the map (skipped): {', '.join(skipped) or 'none'}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
